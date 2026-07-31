from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/konalsmac/Local-Desktop/PRODUCTION-LEVEL-DASES")
BUILD = ROOT / "investor_materials" / "build"
OUT = ROOT / "investor_materials" / "Big_Chalk_Box_Investor_Business_Plan.docx"
LOGO = ROOT / "public" / "logo_new" / "logo.png"

GREEN = "145A38"
DARK_GREEN = "0D4028"
MID_GREEN = "2D6E4C"
GOLD = "BF9B30"
DARK = "1A2421"
INK = "26322E"
SLATE = "5F6E67"
MUTED = "78857F"
CREAM = "FCFCF7"
PALE_GREEN = "EAF1EC"
PALE_GOLD = "F7F3E8"
LIGHT_GRAY = "F3F5F4"
WHITE = "FFFFFF"
RED = "9B1C1C"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        if edge not in edges:
            continue
        edge_data = edges[edge]
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_grid = tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value=True):
    paragraph.paragraph_format.keep_together = value


def add_hyperlink(paragraph, text, url, color=GREEN, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color_node = OxmlElement("w:color")
    color_node.set(qn("w:val"), color)
    rpr.append(color_node)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rpr.append(rfonts)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_picture_alt(paragraph, description, title=None):
    doc_props = paragraph._p.xpath(".//wp:docPr")
    if not doc_props:
        return
    doc_prop = doc_props[-1]
    doc_prop.set("descr", description)
    doc_prop.set("title", title or description)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_separate, cached, fld_char_end])
    set_run_font(run, size=8.5, color=MUTED)


def add_numbering(doc, ordered=False):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "280")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rpr.append(rfonts)
    lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abstract_id))
    num.append(abs_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, bold_lead=None):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    ppr.append(num_pr)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        set_run_font(run, bold=True, color=DARK)
        run2 = p.add_run(text[len(bold_lead):])
        set_run_font(run2, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, color=INK)
    return p


def add_body(doc, text, bold_lead=None, color=INK, italic=False, after=8, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.26
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True, color=DARK)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, color=color, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, color=color, italic=italic)
    return p


def add_kicker(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    set_keep_with_next(p)
    r = p.add_run(text.upper())
    set_run_font(r, size=8.5, color=GOLD, bold=True)
    r.font.letter_spacing = Pt(1.1) if hasattr(r.font, "letter_spacing") else None
    return p


def add_callout(doc, label, title, body, fill=PALE_GOLD, accent=GOLD):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_row_cant_split(table.rows[0])
    set_table_geometry(table, [9360], 120)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, bottom=150, start=190, end=190)
    set_cell_border(
        cell,
        top={"val": "single", "sz": "4", "color": fill},
        bottom={"val": "single", "sz": "4", "color": fill},
        end={"val": "single", "sz": "4", "color": fill},
        start={"val": "single", "sz": "18", "color": accent},
    )
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label.upper())
    set_run_font(r, size=8.2, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    r2 = p2.add_run(title)
    set_run_font(r2, size=13.5, color=DARK, bold=True)
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_after = Pt(0)
    p3.paragraph_format.line_spacing = 1.18
    r3 = p3.add_run(body)
    set_run_font(r3, size=10.3, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_table(doc, headers, rows, widths, header_fill=GREEN, header_color=WHITE,
              font_size=9.2, alignments=None, first_col_bold=False, zebra=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths, 120)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_row_cant_split(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(text))
        set_run_font(run, size=8.7, color=header_color, bold=True)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row_idx, row_values in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        if zebra and row_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, LIGHT_GRAY)
        for col_idx, value in enumerate(row_values):
            cell = cells[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            align = WD_ALIGN_PARAGRAPH.LEFT
            if alignments and col_idx < len(alignments):
                align = alignments[col_idx]
            p.alignment = align
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=INK, bold=(first_col_bold and col_idx == 0))
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    table.rows[-1].cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_source_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, size=8.2, color=MUTED, italic=True)
    return p


def add_section_heading(doc, number, title, subtitle=None):
    add_kicker(doc, f"SECTION {number}")
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(title)
    if subtitle:
        sub = doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(10)
        r = sub.add_run(subtitle)
        set_run_font(r, size=11.5, color=SLATE, italic=True)
    return p


def add_page_break(doc):
    # The cover is the only forced page. Later section boundaries flow naturally,
    # preventing a short spillover paragraph from creating a nearly empty page.
    break_count = getattr(doc, "_bcb_break_count", 0)
    if break_count == 0:
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
    doc._bcb_break_count = break_count + 1


def create_lifecycle_chart(path):
    steps = [
        ("Syllabus", "structure"),
        ("Generate", "question paper"),
        ("Moderate", "quality + coverage"),
        ("Build", "answers + rubric"),
        ("Evaluate", "handwritten scripts"),
        ("Review", "faculty approval"),
        ("Improve", "feedback + analytics"),
    ]
    width, height = 2250, 405
    canvas = Image.new("RGB", (width, height), "#FCFCF7")
    draw = ImageDraw.Draw(canvas)
    regular = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 24)
    bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 31)
    box_w, box_h, gap = 270, 190, 49
    start_x, y = 18, 105
    for i, (title, detail) in enumerate(steps):
        x = start_x + i * (box_w + gap)
        fill = "#EAF1EC" if i not in (2, 4) else "#F7F3E8"
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=26, fill=fill, outline="#B9C8BF", width=3)
        title_box = draw.textbbox((0, 0), title, font=bold)
        detail_box = draw.textbbox((0, 0), detail, font=regular)
        draw.text((x + (box_w - (title_box[2] - title_box[0])) / 2, y + 56), title, font=bold, fill="#1A2421")
        draw.text((x + (box_w - (detail_box[2] - detail_box[0])) / 2, y + 112), detail, font=regular, fill="#5F6E67")
        if i < len(steps) - 1:
            x1, x2, cy = x + box_w + 8, x + box_w + gap - 9, y + box_h // 2
            draw.line((x1, cy, x2, cy), fill="#BF9B30", width=5)
            draw.polygon([(x2, cy), (x2 - 15, cy - 10), (x2 - 15, cy + 10)], fill="#BF9B30")
    canvas.save(path, quality=96)


def create_financial_chart(path):
    years = ["FY27", "FY28", "FY29", "FY30"]
    revenue = [0.25, 1.80, 8.10, 26.00]
    ebitda = [-1.55, -2.18, -0.84, 6.80]
    width, height = 1710, 735
    canvas = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(canvas)
    regular = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 25)
    small = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 22)
    bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 32)
    value_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 23)
    draw.text((105, 32), "Illustrative base case", font=bold, fill="#1A2421")
    draw.rectangle((108, 87, 137, 116), fill="#145A38")
    draw.text((149, 87), "Revenue", font=small, fill="#5F6E67")
    draw.rectangle((298, 87, 327, 116), fill="#C9D4CE")
    draw.text((339, 87), "EBITDA", font=small, fill="#5F6E67")
    left, top, right, bottom = 120, 145, 1630, 625
    y_min, y_max = -4, 28
    def ypix(value):
        return int(bottom - (value - y_min) / (y_max - y_min) * (bottom - top))
    for tick in [-4, 0, 4, 8, 12, 16, 20, 24, 28]:
        y = ypix(tick)
        draw.line((left, y, right, y), fill="#E6EBE8", width=2)
        label = str(tick)
        box = draw.textbbox((0, 0), label, font=small)
        draw.text((left - 18 - (box[2] - box[0]), y - 12), label, font=small, fill="#5F6E67")
    zero_y = ypix(0)
    draw.line((left, zero_y, right, zero_y), fill="#77847E", width=3)
    group_w = (right - left) / 4
    bar_w = 92
    for i, year in enumerate(years):
        center = left + group_w * (i + 0.5)
        for offset, value, color in ((-52, revenue[i], "#145A38"), (52, ebitda[i], "#C9D4CE" if ebitda[i] < 0 else "#BF9B30")):
            x1 = int(center + offset - bar_w / 2)
            x2 = int(center + offset + bar_w / 2)
            vy = ypix(value)
            draw.rectangle((x1, min(zero_y, vy), x2, max(zero_y, vy)), fill=color)
            label = f"{value:g}"
            bbox = draw.textbbox((0, 0), label, font=value_font)
            label_y = vy - 34 if value >= 0 else vy + 8
            draw.text((center + offset - (bbox[2] - bbox[0]) / 2, label_y), label, font=value_font, fill="#145A38" if offset < 0 else "#5F6E67")
        bbox = draw.textbbox((0, 0), year, font=regular)
        draw.text((center - (bbox[2] - bbox[0]) / 2, bottom + 26), year, font=regular, fill="#5F6E67")
    axis_label = "INR crore"
    draw.text((18, 145), axis_label, font=small, fill="#5F6E67")
    canvas.save(path, quality=96)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = rgb(DARK_GREEN)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = rgb(SLATE)
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.paragraph_format.space_after = Pt(8)

    for name, size, color, before, after in (
        ("Heading 1", 16, GREEN, 18, 10),
        ("Heading 2", 13, GREEN, 12, 6),
        ("Heading 3", 11.5, MID_GREEN, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Small Label" not in styles:
        small = styles.add_style("Small Label", WD_STYLE_TYPE.PARAGRAPH)
    else:
        small = styles["Small Label"]
    small.font.name = "Calibri"
    small.font.size = Pt(8.5)
    small.font.bold = True
    small.font.color.rgb = rgb(GOLD)
    small.paragraph_format.space_after = Pt(4)


def configure_section(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    p = first_footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CONFIDENTIAL · DRAFT FOR INVESTOR DISCUSSION")
    set_run_font(r, size=8, color=MUTED, bold=True)

    header = section.header
    hp = header.paragraphs[0]
    hp.clear()
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r1 = hp.add_run("BIG CHALK BOX")
    set_run_font(r1, size=8.2, color=GREEN, bold=True)
    r2 = hp.add_run("\tINVESTOR BUSINESS PLAN · JULY 2026")
    set_run_font(r2, size=8.2, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    left = fp.add_run("CONFIDENTIAL · MANAGEMENT ASSUMPTIONS WHERE NOTED")
    set_run_font(left, size=8, color=MUTED)
    mid = fp.add_run("\tPage ")
    set_run_font(mid, size=8, color=MUTED)
    add_page_field(fp)
    return section


def build_document():
    BUILD.mkdir(parents=True, exist_ok=True)
    create_lifecycle_chart(BUILD / "assessment_lifecycle.png")
    create_financial_chart(BUILD / "financial_base_case.png")

    doc = Document()
    configure_styles(doc)
    configure_section(doc)
    bullet_id = add_numbering(doc, ordered=False)
    number_id = add_numbering(doc, ordered=True)

    # Cover: proposal_centerpiece pattern with branded, restrained treatment.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LOGO), width=Inches(0.78))
    set_picture_alt(p, "Big Chalk Box green chalk-box logo", "Big Chalk Box logo")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("BIG CHALK BOX")
    set_run_font(r, size=28, color=DARK_GREEN, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Investor Business Plan")
    set_run_font(r, size=17, color=DARK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Building the assessment operating system for handwritten education")
    set_run_font(r, size=12.2, color=SLATE, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Prepared 31 July 2026 · India · Confidential")
    set_run_font(r, size=9.5, color=MUTED, bold=True)

    metrics = doc.add_table(rows=1, cols=4)
    set_table_geometry(metrics, [2340, 2340, 2340, 2340], 120)
    set_row_cant_split(metrics.rows[0])
    metric_values = [
        ("400+", "sheets evaluated"),
        ("20+", "educators onboard"),
        ("98%", "rubric accuracy"),
        ("~15 sec", "processing / sheet"),
    ]
    for idx, (value, label) in enumerate(metric_values):
        cell = metrics.cell(0, idx)
        set_cell_shading(cell, PALE_GREEN if idx != 2 else PALE_GOLD)
        set_cell_margins(cell, top=140, bottom=140, start=100, end=100)
        set_cell_border(
            cell,
            top={"val": "single", "sz": "5", "color": "D3DED7"},
            bottom={"val": "single", "sz": "5", "color": "D3DED7"},
            start={"val": "single", "sz": "5", "color": "D3DED7"},
            end={"val": "single", "sz": "5", "color": "D3DED7"},
        )
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(value)
        set_run_font(r1, size=16, color=GREEN, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(label)
        set_run_font(r2, size=8.2, color=SLATE)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(
        doc,
        "Investment proposition",
        "Own the full assessment lifecycle, not one grading feature.",
        "Question Paper Generation and Moderation provide a low-friction entry point. DASES expands the account into recurring, high-value handwritten evaluation, faculty review, student feedback, and institutional analytics.",
        fill=PALE_GOLD,
        accent=GOLD,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    r = p.add_run("Company-reported product metrics; validation plan included in this document.")
    set_run_font(r, size=8, color=MUTED, italic=True)

    add_page_break(doc)

    # 1. Executive summary
    add_section_heading(doc, "01", "Executive Summary", "A focused investment case for an early-stage academic infrastructure company")
    add_callout(
        doc,
        "Recommended seed proposal",
        "₹4.0 crore for approximately 18 months of runway",
        "Use the round to reach 50 paying institutions, a ₹2.0 crore ARR run-rate, 70%+ gross-margin trajectory, independently documented accuracy, and enterprise-grade privacy/security readiness. The amount and milestones are management recommendations, not current commitments.",
        fill=PALE_GREEN,
        accent=GREEN,
    )
    add_body(doc, "Big Chalk Box Pvt. Ltd., founded in 2024, is building AI infrastructure for the most operationally heavy part of education: designing, moderating, evaluating, and learning from handwritten assessments. Its flagship product, DASES (Digital Academic Student Evaluation System), reads handwritten scripts, scores each answer against faculty-defined rubrics, and generates detailed per-question feedback for faculty review.")
    add_body(doc, "The core investment thesis is stronger than a single AI-grading feature. Big Chalk Box can become an assessment workflow system that begins before the exam and compounds after it: syllabus-aligned paper generation, quality moderation, sample answers, marking rubrics, script evaluation, human approval, student feedback, and performance analytics in one institutional record.")
    add_body(doc, "Why now. India registered 64,756 higher-education institutions and nearly 4.5 crore students in AISHE 2023-24. A 2025 FICCI-EY-Parthenon survey reported that 38% of participating Indian HEIs cited automated grading as an AI application, while also emphasizing phased pilots, governance, fairness, and data protection. The market is large, but trust and workflow integration—not generic model access—will determine adoption. [5][6]")

    doc.add_paragraph("The investor case", style="Heading 2")
    for text in [
        "Clear pain with budget value: grading consumes faculty time, result cycles are slow, and feedback is inconsistent or thin.",
        "A wedge-and-expand model: ₹400 + GST per question paper lowers adoption friction; DASES evaluation creates larger, recurring institutional contracts.",
        "Early technical proof: the company reports 400+ sheets evaluated, 20+ educators onboard, 98% rubric accuracy, ~15 seconds per sheet, and 500-sheet parallel capacity. [1][2]",
        "A credible India-first position: handwritten descriptive formats, internal-choice structures, syllabus mapping, Bloom's taxonomy, and price sensitivity are central—not edge cases.",
        "A defensibility path: workflow data, approved rubrics, faculty overrides, integrations, audit trails, and trust can become harder to replace than the underlying AI models.",
    ]:
        add_list_item(doc, text, bullet_id)

    doc.add_paragraph("What must be proven in the next 18 months", style="Heading 2")
    add_table(
        doc,
        ["Proof point", "18-month target", "Investor significance"],
        [
            ["Commercial", "50 paying institutions; ₹2.0 crore ARR run-rate", "Demonstrates repeatable conversion beyond founder-led pilots"],
            ["Retention", ">90% gross logo retention; multi-module expansion", "Shows workflow stickiness and cross-sell"],
            ["Economics", "70%+ gross-margin trajectory; CAC payback <12 months", "Supports venture-scale operating leverage"],
            ["Trust", "Independent benchmark + human-review policy", "De-risks grading accuracy and institutional adoption"],
            ["Governance", "DPDP-ready controls, audit logs, retention/deletion workflows", "Reduces enterprise and regulatory friction"],
        ],
        [1800, 2700, 4860],
        font_size=8.9,
        first_col_bold=True,
        zebra=True,
    )

    add_page_break(doc)

    # 2. Company and problem
    add_section_heading(doc, "02", "Company, Mission & Problem", "Educators should spend time on teaching and judgment—not repetitive paper operations")
    add_body(doc, "Big Chalk Box describes itself as an India-based EdTech engineering company built by engineers and educators. The company's public mission is to automate the tedious, high-effort parts of academic work while keeping educators in control. DASES is the current flagship; Question Paper Generation and Moderation extend the company into a full academic assessment suite. [2][3]")

    doc.add_paragraph("The operational problem", style="Heading 2")
    add_table(
        doc,
        ["Stage", "Current institutional friction", "Business consequence"],
        [
            ["Paper design", "Questions are assembled manually across syllabi, outcomes, difficulty levels, and prior papers", "Slow cycles; uneven quality; repeated questions"],
            ["Moderation", "Review depends on scarce senior faculty and inconsistent checklists", "Ambiguity, syllabus gaps, appeals, and rework"],
            ["Rubric creation", "Model answers and criteria are often late, coarse, or inconsistent", "Evaluator drift and weak partial-credit logic"],
            ["Evaluation", "Hundreds of handwritten scripts are reviewed sequentially", "Delayed results and faculty workload spikes"],
            ["Feedback", "Time pressure reduces feedback to a score or short remark", "Students do not know how to improve"],
            ["Analytics", "Question-level evidence is trapped in paper files", "Little insight into topic gaps or assessment quality"],
        ],
        [1450, 4300, 3610],
        font_size=8.7,
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Why existing approaches leave a gap", style="Heading 2")
    for text in [
        "Manual grading preserves judgment but does not scale consistently across large batches.",
        "On-screen marking digitizes handling, but many workflows still require human scoring for every response.",
        "Global grading tools digitize handwritten submissions and may group similar answers, but are not necessarily designed for autonomous rubric-based evaluation of Indian descriptive exam formats. [7]",
        "Generic AI tools can draft questions or score text, but lack institutional controls, original-script traceability, audit trails, and an end-to-end academic workflow.",
    ]:
        add_list_item(doc, text, bullet_id)

    add_callout(
        doc,
        "Design principle",
        "AI proposes and processes; faculty remain accountable.",
        "The defensible operating model is human-in-the-loop: faculty approve rubrics, review low-confidence or high-stakes answers, override scores, and publish results. This positions Big Chalk Box as decision-support infrastructure rather than an unreviewable black-box grader.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    add_page_break(doc)

    # 3. Product suite
    add_section_heading(doc, "03", "Product Suite & Workflow", "One assessment record from syllabus to student feedback")
    add_body(doc, "The suite is organized around a single lifecycle. Each module can sell independently, but the strategic value comes from shared syllabus, question, answer, rubric, and evaluation data flowing forward without re-entry.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(str(BUILD / "assessment_lifecycle.png"), width=Inches(6.35))
    set_picture_alt(p, "Assessment lifecycle: syllabus, generate, moderate, build answers and rubric, evaluate, faculty review, feedback and analytics", "Big Chalk Box assessment lifecycle")
    add_source_note(doc, "Figure 1. Target assessment lifecycle. Current module status and roadmap should be confirmed before external circulation.")

    add_table(
        doc,
        ["Module", "Customer job", "Current value proposition"],
        [
            ["Question Paper Moderation", "Audit a draft before approval", "Syllabus coverage, difficulty/Bloom's balance, ambiguity, duplication, structure, and suggested revisions"],
            ["Question Paper Generation", "Create a compliant paper quickly", "Syllabus-aligned generation by section, difficulty, Bloom's targets, sample answers, rubrics, and print-ready output"],
            ["DASES Evaluation", "Evaluate handwritten scripts at scale", "Handwriting recognition, rubric-based scoring, per-question feedback, faculty review, and student reports"],
            ["Analytics layer", "Improve teaching and assessment quality", "Coverage, question performance, criterion gaps, and—in later releases—outcome attainment"],
        ],
        [1900, 2400, 5060],
        font_size=8.8,
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Product edge that matters to institutions", style="Heading 2")
    for text in [
        "Institution-specific configuration: syllabi, course outcomes, formats, internal choices, mark allocations, and rubric rules.",
        "Original-to-decision traceability: faculty should be able to inspect the scan, extracted content, criterion score, rationale, override, and approver.",
        "Pre-exam and post-exam continuity: the same approved question paper, sample answer, and rubric flow into evaluation.",
        "Operational scale: company-reported processing of up to 500 sheets in parallel at approximately 15 seconds per sheet. [1][3]",
        "Feedback as a product: every student can receive question-level, criterion-level guidance—not merely a total score.",
    ]:
        add_list_item(doc, text, bullet_id)

    doc.add_paragraph("September 2026 release commitment", style="Heading 2")
    add_body(doc, "The supplied pricing roadmap states that Sample Answers and Marking Rubrics for Question Paper Moderation are scheduled for release in September 2026. This release is strategically important because it closes the handoff from paper quality to DASES evaluation and increases cross-module attach rate.")

    add_page_break(doc)

    # 4. Business model
    add_section_heading(doc, "04", "Business Model & Pricing", "Preserve the low-friction entry point; convert usage into committed institutional revenue")
    add_table(
        doc,
        ["Offering", "Public price", "Commercial role"],
        [
            ["Question Paper Moderation", "₹400 + GST / paper", "Low-risk paid wedge; creates evidence and an evaluation-ready rubric"],
            ["Question Paper Generation", "₹400 + GST / paper", "Faculty productivity wedge; generates sample answers and marking rubrics"],
            ["DASES Evaluation", "Custom pricing", "Primary expansion product; price by volume, modules, support, and integration needs"],
        ],
        [2400, 1800, 5160],
        font_size=9,
        first_col_bold=True,
        zebra=True,
    )
    add_source_note(doc, "Source: management-supplied Pricing — BigChalkBox Academic Suite, July 2026. GST is excluded from revenue calculations.")

    doc.add_paragraph("Recommended packaging architecture", style="Heading 2")
    add_table(
        doc,
        ["Commercial layer", "Illustrative structure", "Purpose"],
        [
            ["Pay-as-you-go", "Retain ₹400/paper for generation and moderation", "Fast departmental adoption; transparent value demonstration"],
            ["Department commitment", "₹2.4–4.8 lakh annual minimum; shared credits + evaluation volume", "Creates recurring revenue and predictable usage"],
            ["Institution contract", "₹8–20 lakh annual minimum; multi-department controls, integrations, SLA, analytics", "Raises ACV and embeds the platform institution-wide"],
            ["Overages / services", "Per-sheet evaluation overage; implementation/integration fees", "Protects gross margin and funds non-standard work"],
        ],
        [2000, 3550, 3810],
        font_size=8.8,
        first_col_bold=True,
        zebra=True,
    )
    add_source_note(doc, "Illustrative management recommendation—not current published pricing. Final tiers require unit-cost, willingness-to-pay, and procurement testing.")

    doc.add_paragraph("Economic logic", style="Heading 2")
    economic_number_id = add_numbering(doc, ordered=True)
    for text in [
        "Land: one department buys generation or moderation for a small number of papers.",
        "Prove: a structured pilot measures faculty hours, agreement with approved scores, review time, and student feedback quality.",
        "Expand: evaluation volume, additional departments, student reports, analytics, and integrations increase annual contract value.",
        "Renew: reusable syllabi, rubrics, workflows, permissions, and reporting history reduce the cost of each subsequent exam cycle.",
    ]:
        add_list_item(doc, text, economic_number_id)

    add_callout(
        doc,
        "Pilot ROI instrument",
        "Measure value before negotiating a campus-wide contract.",
        "For every pilot, record baseline faculty hours, AI processing time, faculty review time, score agreement, override rate, result turnaround, feedback coverage, compute cost, and willingness to renew. The investor-quality metric is not 'sheets processed'; it is repeatable paid conversion with improving gross margin.",
        fill=PALE_GREEN,
        accent=GREEN,
    )

    add_page_break(doc)

    # 5. Traction
    add_section_heading(doc, "05", "Traction & Evidence", "Promising product proof; the next round must convert it into commercial proof")
    add_table(
        doc,
        ["Metric", "Reported result", "Diligence standard for next round"],
        [
            ["Sheets evaluated", "400+", "Batch-level log with subjects, institutions, dates, and review outcomes"],
            ["Educators onboard", "20+", "Active vs. trial users; institution count; 30/90-day retention"],
            ["Rubric accuracy", "98%", "Predefined benchmark, sample size, blind human comparison, confidence interval"],
            ["Processing time", "~15 sec/sheet", "p50/p95 latency by page count, subject, image quality, and batch size"],
            ["Parallel scale", "Up to 500 sheets", "Load test, failure/retry rate, queue behavior, and cost per sheet"],
            ["Program support", "Runway Incubator; Google for Startups; AWS Startups listed", "Acceptance/credit documentation and active benefit status"],
        ],
        [1900, 1800, 5660],
        font_size=8.8,
        first_col_bold=True,
        zebra=True,
    )
    add_source_note(doc, "All results in this table are company-reported on DASES public pages and should be substantiated in the investor data room. [1][2]")

    doc.add_paragraph("Evidence hierarchy investors will reward", style="Heading 2")
    for text in [
        "Best: paid renewals and expansion across departments.",
        "Strong: signed annual contracts, purchase orders, or paid pilots with referenceable institutions.",
        "Useful: dated LOIs with scope, expected contract value, and procurement owner.",
        "Supporting: educator usage, testimonial, benchmark, or demo engagement.",
        "Weak alone: waitlist size, page views, unqualified signups, or unverified accuracy claims.",
    ]:
        add_list_item(doc, text, bullet_id)

    add_callout(
        doc,
        "Messaging discipline",
        "Convert technical claims into auditable institutional outcomes.",
        "Use 'company-reported' until an independent or jointly signed validation exists. Replace broad claims such as 'human-level accuracy' with benchmark definitions, exceptions, and the faculty review policy. Credibility improves the deal more than an inflated metric.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    add_page_break(doc)

    # 6. Market
    add_section_heading(doc, "06", "Market Opportunity", "A large institutional base with a narrow, measurable beachhead")
    add_body(doc, "AISHE 2023-24 registered 1,289 universities, 48,246 colleges, and 15,221 standalone institutions—64,756 HEIs in total—with nearly 4.5 crore students and 17.3 lakh teachers. Seventy percent of responding colleges were private unaided, an attractive segment for faster procurement and measurable operating ROI. [5]")

    add_table(
        doc,
        ["Layer", "Scope assumption", "Illustrative ACV", "Annual value"],
        [
            ["TAM", "64,756 registered Indian HEIs", "₹6 lakh", "₹3,885 crore"],
            ["SAM", "20,000 priority private/autonomous/technical institutions", "₹6 lakh", "₹1,200 crore"],
            ["4-year SOM", "450 contracted institutions", "₹8 lakh", "₹36 crore ACV"],
        ],
        [1200, 3720, 1800, 2640],
        font_size=9,
        first_col_bold=True,
        zebra=True,
    )
    add_source_note(doc, "Bottom-up management sizing, not a third-party market forecast. TAM/SAM use an illustrative blended annual contract value; SOM is an execution target. Recognized FY30 revenue is lower than exit ACV because customers ramp through the year.")

    doc.add_paragraph("Beachhead customer profile", style="Heading 2")
    for text in [
        "Private and autonomous universities or colleges with 1,000–10,000 students.",
        "Engineering, management, science, and professional programs with frequent descriptive assessments.",
        "A Controller of Examinations, Dean, HOD, or quality/accreditation owner with a visible result-turnaround problem.",
        "Existing scanning capability or willingness to adopt a standardized scan workflow.",
        "A department prepared to run a faculty-reviewed pilot on 100–500 scripts.",
    ]:
        add_list_item(doc, text, bullet_id)

    doc.add_paragraph("Timing signal", style="Heading 2")
    add_body(doc, "The 2025 FICCI-EY-Parthenon survey of 30 Indian HEIs found that 57% had an institutional AI policy and another 40% were developing one; 38% cited automated grading among academic applications. The same research calls for phased adoption and highlights fairness, transparency, accountability, and data protection. That combination favors a pilot-first vendor with strong governance rather than a generic AI tool. [6]")

    add_page_break(doc)

    # 7. GTM
    add_section_heading(doc, "07", "Go-to-Market Strategy", "A reference-led institutional motion with productized pilots")
    add_table(
        doc,
        ["Stage", "Offer", "Decision owner", "Target evidence"],
        [
            ["1. Diagnose", "Exam workflow assessment + live demo", "HOD / Dean / COE", "Qualified pain, volume, budget, timeline"],
            ["2. Pilot", "1 department; 1 paper; 100–500 scripts", "Faculty champion + evaluator", "Accuracy, review time, turnaround, ROI"],
            ["3. Convert", "Annual department commitment", "Dean / procurement", "Paid contract and recurring calendar"],
            ["4. Expand", "More departments + modules + integrations", "COE / CIO / academic leadership", "Higher ACV and multi-threaded use"],
            ["5. Reference", "Case study + peer introduction", "Institution sponsor", "Lower CAC and regional trust"],
        ],
        [1200, 2800, 1900, 3460],
        font_size=8.7,
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Channel priorities", style="Heading 2")
    for text in [
        "Founder-led direct sales for the first 25–30 institutions to refine the playbook.",
        "Regional academic networks, incubators, and educator champions for references.",
        "Scanner/document-digitization partners that already serve exam cells.",
        "ERP/LMS and accreditation consultants as implementation and referral partners.",
        "Selective coaching-institute and school pilots only where the same workflow and unit economics apply.",
    ]:
        add_list_item(doc, text, bullet_id)

    doc.add_paragraph("Operating funnel targets", style="Heading 2")
    add_table(
        doc,
        ["KPI", "Initial target", "Why it matters"],
        [
            ["Qualified demo → pilot", ">35%", "Indicates the pain and product promise are real"],
            ["Pilot → paid annual", ">40%", "Tests willingness to pay and procurement readiness"],
            ["Time to first evaluated batch", "<14 days", "Reduces implementation friction"],
            ["Gross logo retention", ">90%", "Validates recurring exam-cycle value"],
            ["Net revenue retention", ">115%", "Shows module and department expansion"],
            ["CAC payback", "<12 months", "Keeps institutional sales capital-efficient"],
            ["Receivable days", "<75 days", "Prevents procurement cycles from consuming runway"],
        ],
        [2500, 1800, 5060],
        font_size=8.8,
        first_col_bold=True,
        zebra=True,
    )
    add_source_note(doc, "Targets are management assumptions for operating discipline; they are not historical performance.")

    add_page_break(doc)

    # 8. Competition
    add_section_heading(doc, "08", "Competitive Landscape & Defensibility", "The moat must be workflow trust and distribution—not access to an AI model")
    add_table(
        doc,
        ["Alternative", "What it does well", "Gap / strategic response"],
        [
            ["Gradescope / Turnitin", "Digitizes paper assessments, supports handwritten work, dynamic rubrics, AI-assisted answer grouping, analytics, LMS workflows", "Differentiate on India-specific descriptive automation, pre-exam generation/moderation, local formats, and institutional economics. [7]"],
            ["Eklavvya", "Broad exam platform with question generation, handwritten AI evaluation, multilingual workflows, identity masking, and audit/re-evaluation capabilities", "Closest direct competitor. Win on product coherence, measurable evaluation quality, faster pilots, transparent pricing wedge, and focused assessment lifecycle. [8]"],
            ["On-screen marking vendors", "Mature scan, allocation, blind marking, and audit workflows", "Integrate where useful; lead with AI-assisted scoring and feedback economics"],
            ["Generic AI / OCR", "Low-cost drafting, OCR, and experimentation", "Emphasize original-script traceability, permissions, review, rubrics, data controls, and repeatable institutional operations"],
            ["Manual process", "Trusted judgment and flexible exception handling", "Retain faculty authority while automating repetitive work"],
        ],
        [1800, 3900, 3660],
        font_size=8.4,
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Defensibility roadmap", style="Heading 2")
    for text in [
        "Workflow depth: one canonical assessment record from syllabus through approved student result.",
        "Institutional configuration: paper formats, outcomes, rubrics, policies, access controls, and escalation rules.",
        "Quality system: benchmark sets, confidence thresholds, faculty overrides, drift monitoring, and appeal outcomes.",
        "Integration surface: scanners, rosters, ERP/LMS gradebooks, identity, and institutional reporting.",
        "Distribution: reference institutions, evaluator champions, partners, and implementation know-how by segment.",
        "Privacy-preserving learning: improve prompts, workflows, and evaluation logic only under clear contractual data rights; do not assume student data can be reused for model training.",
    ]:
        add_list_item(doc, text, bullet_id)

    add_callout(
        doc,
        "Investor reality check",
        "Accuracy claims are necessary; reliable operations create the moat.",
        "Vision and language models will continue improving and commoditizing. Durable value comes from evaluation governance, approved workflows, data rights, integrations, evidence, and trusted distribution inside institutions.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    add_page_break(doc)

    # 9. Technology and governance
    add_section_heading(doc, "09", "Technology, Quality & Governance", "Build for a challenged grade, not only a successful demo")
    doc.add_paragraph("Target evaluation control loop", style="Heading 2")
    control_number_id = add_numbering(doc, ordered=True)
    for text in [
        "Ingest: validate scan quality, identity mapping, page count, and file integrity.",
        "Recognize: extract handwriting and preserve the original-image reference.",
        "Map: link every response to the correct question, including internal choices and continuation pages.",
        "Evaluate: score each rubric criterion, produce rationale, and record model/version/configuration.",
        "Escalate: route low-confidence, anomalous, or high-stakes responses to faculty.",
        "Approve: allow faculty override with reason codes; preserve the full audit trail.",
        "Publish: generate institutional reports and apply retention/deletion rules.",
    ]:
        add_list_item(doc, text, control_number_id)

    doc.add_paragraph("Minimum enterprise control set", style="Heading 2")
    add_table(
        doc,
        ["Control area", "Required capability"],
        [
            ["Accuracy", "Subject- and question-type benchmark; inter-rater protocol; override-rate tracking; confidence thresholds"],
            ["Fairness", "Blind review where appropriate; cohort error analysis; handwriting-quality and language-segment checks"],
            ["Auditability", "Immutable event trail for upload, extraction, score, override, approver, and publication"],
            ["Privacy", "Data inventory, lawful basis/notice support, purpose limitation, retention schedule, deletion/export workflow, vendor DPAs"],
            ["Security", "Encryption, RBAC, least privilege, isolated tenants, secrets management, backups, incident response, vulnerability testing"],
            ["Reliability", "Retry logic, queue visibility, idempotency, model fallback policy, service-level monitoring, disaster recovery"],
            ["Academic control", "Faculty-approved rubrics, explicit appeal/re-evaluation process, no autonomous publication by default"],
        ],
        [2000, 7360],
        font_size=8.7,
        first_col_bold=True,
        zebra=True,
    )

    add_body(doc, "India's Digital Personal Data Protection Rules, 2025 were notified in November 2025 with a staged enforcement timeline for the Act and Rules. Because answer sheets and student identifiers are personal data, privacy readiness should be part of the product roadmap and contract architecture, with qualified Indian legal counsel confirming obligations and timelines. [9]")

    add_callout(
        doc,
        "Non-negotiable policy",
        "Faculty approval remains the final grading authority.",
        "At this stage, DASES should default to assistive evaluation with visible rationale, editable scores, confidence-based escalation, and a documented re-evaluation route. This reduces both adoption risk and academic harm.",
        fill=PALE_GREEN,
        accent=GREEN,
    )

    add_page_break(doc)

    # 10. Roadmap
    add_section_heading(doc, "10", "18–36 Month Execution Roadmap", "Fund milestones that change the next financing outcome")
    add_table(
        doc,
        ["Period", "Product & trust", "Commercial", "Organization"],
        [
            ["0–6 months", "Ship Sept 2026 sample answers/rubrics; benchmark protocol; confidence routing; cost telemetry; core DPDP controls", "10–15 paying institutions; 3 referenceable pilots; productized pilot kit", "Strengthen AI/ML, academic QA, and customer success"],
            ["6–12 months", "Multi-department administration; audit/re-evaluation; ERP/LMS/roster connectors; usage billing", "25–30 paying institutions; repeatable regional funnel; first expansion contracts", "Hire enterprise sales and implementation leads"],
            ["12–18 months", "Independent validation; security audit/certification readiness; analytics v1", "50 paying institutions; ₹2 crore ARR run-rate; >90% gross retention", "Board reporting, finance controls, and second-line leadership"],
            ["18–36 months", "Language/subject expansion based on demand; accreditation analytics; partner APIs", "200 institutions by FY29; selective school/coaching channels; 2–3 regional clusters", "Scale GTM, support, security, and model operations"],
        ],
        [1300, 3020, 3000, 2040],
        font_size=8.1,
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Hiring plan by capability", style="Heading 2")
    add_table(
        doc,
        ["Capability", "Near-term priority", "Why now"],
        [
            ["AI/ML & evaluation quality", "Senior owner + evaluation engineer", "Benchmarking, confidence, cost, model reliability"],
            ["Academic quality", "Advisor panel + full-time program lead", "Rubric design, subject coverage, faculty trust"],
            ["Product/engineering", "2–3 full-stack/platform hires", "Admin, integrations, billing, reliability"],
            ["Enterprise sales", "2 focused sellers after playbook validation", "Convert pilots without scaling founder dependency too early"],
            ["Customer success", "Implementation lead + support", "Faster first value, stronger retention and references"],
            ["Security/privacy", "Fractional specialist → dedicated owner", "Enterprise diligence and DPDP readiness"],
        ],
        [2200, 3000, 4160],
        font_size=8.7,
        first_col_bold=True,
        zebra=True,
    )

    add_body(doc, "Team disclosure for investor circulation. Add the founders' names, roles, ownership, full-time status, prior technical/education experience, and one proof point per person to the accompanying pitch deck and data room. The public site currently describes the team only as engineers and educators based in India; investors will require named biographies. [2]")

    add_page_break(doc)

    # 11. Financial plan
    add_section_heading(doc, "11", "Illustrative Financial Plan", "A milestone model to test the business—not a substitute for actual historicals")
    add_callout(
        doc,
        "Model status",
        "Base-case management assumptions; replace with actuals before circulation.",
        "The model below is intentionally bottom-up and conservative on early revenue recognition. It assumes annual commitments grow as accounts add evaluation volume, departments, analytics, and integrations. It excludes GST and does not include financing proceeds as revenue.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(BUILD / "financial_base_case.png"), width=Inches(6.25))
    set_picture_alt(p, "Illustrative base-case chart showing revenue rising from 0.25 crore in FY27 to 26 crore in FY30 and EBITDA becoming positive in FY30", "Illustrative financial base case")

    add_table(
        doc,
        ["Metric", "FY27", "FY28", "FY29", "FY30"],
        [
            ["Exit paying institutions", "20", "75", "200", "450"],
            ["Average active institutions", "10", "45", "135", "325"],
            ["Blended revenue / active institution", "₹2.5L", "₹4.0L", "₹6.0L", "₹8.0L"],
            ["Revenue", "₹0.25Cr", "₹1.80Cr", "₹8.10Cr", "₹26.00Cr"],
            ["Gross margin", "58%", "68%", "76%", "80%"],
            ["Gross profit", "₹0.15Cr", "₹1.22Cr", "₹6.16Cr", "₹20.80Cr"],
            ["Operating expenses", "₹1.70Cr", "₹3.40Cr", "₹7.00Cr", "₹14.00Cr"],
            ["EBITDA", "(₹1.55Cr)", "(₹2.18Cr)", "(₹0.84Cr)", "₹6.80Cr"],
        ],
        [2880, 1620, 1620, 1620, 1620],
        font_size=8.7,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Model drivers to validate immediately", style="Heading 2")
    for text in [
        "DASES price realization per sheet, paper, department, and institution.",
        "Model/API, OCR, storage, report generation, and human QA cost per sheet.",
        "Pilot support hours and implementation cost by institution size.",
        "Procurement duration, discounting, payment schedule, and receivable days.",
        "Pilot conversion, renewal, module attach, and department expansion rates.",
        "Headcount plan, founder compensation, cloud credits, and tax/compliance costs.",
    ]:
        add_list_item(doc, text, bullet_id)

    add_source_note(doc, "FY27 means the financial year ending March 2027. Because this plan is dated July 2026, FY27 is a partial execution year. Forecasts should be re-cut using the latest monthly actuals and qualified pipeline.")

    add_page_break(doc)

    # 12. Funding and use of funds
    add_section_heading(doc, "12", "Funding Proposal & Use of Funds", "A seed round sized to evidence, not vanity")
    add_callout(
        doc,
        "Proposed raise",
        "₹4.0 crore primary capital · approximately 18 months",
        "The financing should be positioned as a milestone round to establish repeatable institutional sales, trusted evaluation quality, enterprise readiness, and a ₹2.0 crore ARR run-rate. Final valuation, instrument, and rights require current cap-table, revenue, pipeline, and qualified legal/tax advice.",
        fill=PALE_GREEN,
        accent=GREEN,
    )

    add_table(
        doc,
        ["Use", "%", "Amount", "Expected output"],
        [
            ["Product, AI & engineering", "35%", "₹1.40Cr", "Evaluation quality, workflow depth, reliability, billing, and integrations"],
            ["Go-to-market", "25%", "₹1.00Cr", "Productized pilots, regional pipeline, sales capacity, references"],
            ["Security, privacy & QA", "15%", "₹0.60Cr", "DPDP program, audit trail, benchmark, external testing/certification readiness"],
            ["Cloud & data operations", "12.5%", "₹0.50Cr", "Processing capacity, observability, backups, cost optimization"],
            ["Customer success & implementation", "7.5%", "₹0.30Cr", "Faster onboarding, training, support, institutional deployment"],
            ["Legal, finance & contingency", "5%", "₹0.20Cr", "Contracts, finance controls, IP/data documentation, buffer"],
        ],
        [2500, 1000, 1400, 4460],
        font_size=8.6,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Round milestones", style="Heading 2")
    add_table(
        doc,
        ["By month", "Milestone", "Financing implication"],
        [
            ["6", "10–15 paying institutions; September features live; benchmark protocol running", "Confirms product completion and paid wedge"],
            ["12", "25–30 institutions; >₹1 crore ARR run-rate; 3 referenceable case studies", "Supports repeatability and reference-led growth"],
            ["18", "50 institutions; ₹2 crore ARR run-rate; 70%+ GM trajectory; >90% gross retention", "Creates a credible Series A / extension decision point"],
        ],
        [1400, 4900, 3060],
        font_size=8.8,
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Investor fit", style="Heading 2")
    for text in [
        "Seed investor with B2B SaaS and India institutional-sales experience.",
        "Education, assessment, ERP/LMS, or public-sector networks that accelerate references and procurement.",
        "AI governance, data protection, and enterprise security operating support.",
        "Follow-on capacity for a later round if commercial evidence compounds.",
        "No commercial exclusivity that restricts the company from serving broad institutional segments.",
    ]:
        add_list_item(doc, text, bullet_id)

    add_page_break(doc)

    # 13. Risks
    add_section_heading(doc, "13", "Key Risks & Mitigations", "The plan becomes investable when risks are measurable and owned")
    add_table(
        doc,
        ["Risk", "Potential impact", "Mitigation and leading indicator"],
        [
            ["Accuracy / grading harm", "Loss of trust, disputes, academic harm", "Faculty approval; confidence routing; benchmark by subject/question type; override and appeal tracking"],
            ["Privacy / DPDP", "Contract loss, remediation cost, regulatory exposure", "Data inventory, notice/contract support, retention/deletion, DPAs, incident response, legal review"],
            ["Long procurement cycles", "Slow revenue and high working capital", "Paid department pilots, procurement map, milestone billing, receivable target, channel partners"],
            ["Model and compute costs", "Weak gross margin at scale", "Per-sheet telemetry, model routing, caching, batch processing, overage floors, quarterly price review"],
            ["Competitive response", "Price pressure and feature parity", "Integrated workflow, evidence, references, India-specific formats, fast implementation, integrations"],
            ["Unsubstantiated claims", "Investor/customer credibility loss", "Claim register, methodology appendix, independent validation, signed customer case studies"],
            ["Seasonality", "Revenue concentration around exam cycles", "Annual commitments, multiple modules, multiple institution segments, implementation calendars"],
            ["Founder / key-person risk", "Sales and product bottleneck", "Named owners, second-line hires, board cadence, documentation, delegated customer success"],
        ],
        [1900, 2660, 4800],
        font_size=8.3,
        first_col_bold=True,
        zebra=True,
    )

    add_callout(
        doc,
        "Decision rule",
        "Do not scale autonomous scoring faster than the evidence system.",
        "Expansion should follow demonstrated agreement, explainability, review efficiency, and appeal outcomes by subject and answer type. When confidence is low, route to a person—the platform's credibility is more valuable than marginal automation.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    add_page_break(doc)

    # 14. Diligence readiness
    add_section_heading(doc, "14", "Investor Diligence Readiness", "The fastest route to a better deal is removing uncertainty before term-sheet discussion")
    add_table(
        doc,
        ["Data-room folder", "Minimum content before investor circulation"],
        [
            ["Corporate", "Certificate of incorporation, constitutional documents, cap table, past issuances, board/shareholder approvals, statutory filings"],
            ["Team", "Named founder bios, employment/consulting agreements, ESOP plan, hiring plan, reference checks"],
            ["Commercial", "Revenue by customer/month/module, contracts, invoices, pipeline with stage/value/date, LOIs/POs, churn/renewal"],
            ["Product & traction", "Usage logs, institution/educator definitions, pilot reports, product roadmap, support metrics"],
            ["Accuracy & QA", "Benchmark methodology, dataset composition, human-rater protocol, error analysis, overrides, exceptions, claim register"],
            ["Unit economics", "API/model/OCR/storage cost, gross margin by product, support hours, CAC, sales cycle, receivable days"],
            ["Technology & IP", "Architecture, repositories, access control, IP assignments, third-party licenses, model/vendor terms, disaster recovery"],
            ["Privacy & security", "Data map, privacy notices, customer DPA, subprocessors, retention/deletion, vulnerability tests, incidents"],
            ["Finance & tax", "Monthly P&L, balance sheet, cash flow, bank statements, GST/TDS, budgets, liabilities, related-party transactions"],
        ],
        [2350, 7010],
        font_size=8.4,
        first_col_bold=True,
        zebra=True,
    )

    doc.add_paragraph("Five facts to personalize before sending this plan", style="Heading 2")
    facts_number_id = add_numbering(doc, ordered=True)
    for text in [
        "Founder names, roles, biographies, full-time status, and ownership.",
        "Actual revenue/ARR, cash balance, monthly burn, liabilities, and prior funding.",
        "Paying institutions, named pilots/LOIs, contract values, and qualified pipeline.",
        "DASES pricing, cloud/model cost per sheet, gross margin, and faculty review effort.",
        "Exact accuracy methodology and independent/customer validation evidence.",
    ]:
        add_list_item(doc, text, facts_number_id)

    add_callout(
        doc,
        "Fundraising discipline",
        "Lead with evidence; negotiate dilution after the milestone story is credible.",
        "A strong process creates leverage: clean data room, consistent metrics, multiple interested investors, a defined close date, and a clear view of the minimum capital needed to reach the next value inflection. Avoid presenting an unsupported valuation before the actual commercial evidence is assembled.",
        fill=PALE_GREEN,
        accent=GREEN,
    )

    add_page_break(doc)

    # 15. Sources
    add_section_heading(doc, "15", "Sources & Research Notes", "Research cutoff: 31 July 2026")
    sources = [
        ("[1] DASES homepage", "Company-reported capabilities, performance, affiliations, and contact information.", "https://www.dasesai.com/"),
        ("[2] DASES About", "Company-reported traction, mission, India focus, and founding-team description.", "https://www.dasesai.com/about"),
        ("[3] DASES Solutions", "Product modules and end-to-end assessment workflow.", "https://www.dasesai.com/solutions"),
        ("[4] BigChalkBox pricing", "Management-supplied pricing: ₹400 + GST per generated/moderated paper; custom DASES evaluation pricing.", "https://bigchalkbox.com/pricing"),
        ("[5] Ministry of Education, AISHE 2023-24", "Official HEI, enrolment, teacher, and institution-management statistics.", "https://www.dohe-education.gov.in/static/uploads/2026/07/8616f33dfee644ab6b87a2bd0658b18d.pdf"),
        ("[6] EY India, AI in higher education", "FICCI-EY-Parthenon 2025 survey findings and governance recommendations.", "https://www.ey.com/en_in/insights/education/harnessing-ai-in-higher-education-opportunities-and-the-road-ahead"),
        ("[7] Turnitin Gradescope", "Official description of handwritten support, dynamic rubrics, AI-assisted grouping, and institutional workflow.", "https://in.turnitin.com/products/gradescope/"),
        ("[8] Eklavvya", "Official description of AI question generation and handwritten answer-sheet evaluation capabilities.", "https://www.eklavvya.com/"),
        ("[9] MeitY DPDP Rules 2025", "Official rules, corrigendum, and staged enforcement materials.", "https://www.meity.gov.in/documents/act-and-policies/digital-personal-data-protection-rules-2025-gDOxUjMtQWa"),
    ]
    for title, note, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_together = True
        r = p.add_run(title + ". ")
        set_run_font(r, size=9.4, color=DARK, bold=True)
        r2 = p.add_run(note + " ")
        set_run_font(r2, size=9.4, color=INK)
        add_hyperlink(p, "Open source", url)

    doc.add_paragraph("Notes on interpretation", style="Heading 2")
    for text in [
        "Company metrics and roadmap statements are not independently verified unless explicitly stated.",
        "Market sizing is a bottom-up scenario derived from official institution counts and illustrative annual contract values; it is not a market-research forecast.",
        "Financial projections, packaging, funding amount, and milestone targets are management recommendations created for planning and must be replaced or approved using actual company data.",
        "Competitive descriptions use public vendor materials and should be validated during commercial diligence.",
        "Privacy and fundraising references are strategic planning considerations, not legal, tax, accounting, or investment advice.",
    ]:
        add_list_item(doc, text, bullet_id)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("BIG CHALK BOX PVT. LTD.")
    set_run_font(r, size=12, color=GREEN, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(0)
    add_hyperlink(p2, "bigchalkbox.com", "https://bigchalkbox.com/")
    r2 = p2.add_run("  ·  ")
    set_run_font(r2, size=9, color=MUTED)
    add_hyperlink(p2, "dasesai.com", "https://www.dasesai.com/")

    # Core properties and save.
    doc.core_properties.title = "Big Chalk Box Investor Business Plan"
    doc.core_properties.subject = "Investor-ready business plan and seed financing framework"
    doc.core_properties.author = "Big Chalk Box Pvt. Ltd."
    doc.core_properties.keywords = "Big Chalk Box, DASES, investor business plan, EdTech, AI assessment"
    doc.core_properties.comments = "Prepared from company-provided pricing, public company materials, and cited external research."

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
